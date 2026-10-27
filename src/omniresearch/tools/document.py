"""Document parser for enterprise local documents: PDF, DOCX, CSV, TXT, MD."""

import os
import csv
from typing import List
from omniresearch.core.state import SourceDocument


def parse_pdf(file_path: str, start_id: int = 1) -> List[SourceDocument]:
    """Extract text from PDF pages with citation tracking."""
    documents: List[SourceDocument] = []
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(file_path)
        base_name = os.path.basename(file_path)
        cur_id = start_id
        
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text().strip()
            if text:
                documents.append(
                    SourceDocument(
                        id=cur_id,
                        title=f"{base_name} (Page {page_num + 1})",
                        url_or_path=file_path,
                        snippet=text[:1500],  # Chunk per page
                        source_type="pdf",
                        relevance_score=1.0,
                    )
                )
                cur_id += 1
        doc.close()
    except Exception as exc:
        print(f"[OmniResearch Doc Warning] Error reading PDF {file_path}: {exc}")
    return documents


def parse_docx(file_path: str, start_id: int = 1) -> List[SourceDocument]:
    """Extract text from Word DOCX documents."""
    documents: List[SourceDocument] = []
    try:
        from docx import Document
        doc = Document(file_path)
        base_name = os.path.basename(file_path)
        full_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text.strip())
        
        joined = "\n\n".join(full_text)
        if joined:
            documents.append(
                SourceDocument(
                    id=start_id,
                    title=f"{base_name} (Full Document)",
                    url_or_path=file_path,
                    snippet=joined[:3000],
                    source_type="docx",
                    relevance_score=1.0,
                )
            )
    except Exception as exc:
        print(f"[OmniResearch Doc Warning] Error reading DOCX {file_path}: {exc}")
    return documents


def parse_csv(file_path: str, start_id: int = 1) -> List[SourceDocument]:
    """Extract summary rows from CSV files."""
    documents: List[SourceDocument] = []
    try:
        base_name = os.path.basename(file_path)
        rows_summary = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            headers = next(reader, None)
            if headers:
                rows_summary.append("Columns: " + ", ".join(headers))
            for i, row in enumerate(reader):
                if i < 20:  # First 20 rows sample
                    rows_summary.append(" | ".join(row))
                else:
                    break
        
        if rows_summary:
            documents.append(
                SourceDocument(
                    id=start_id,
                    title=f"{base_name} (Data Summary)",
                    url_or_path=file_path,
                    snippet="\n".join(rows_summary),
                    source_type="csv",
                    relevance_score=1.0,
                )
            )
    except Exception as exc:
        print(f"[OmniResearch Doc Warning] Error reading CSV {file_path}: {exc}")
    return documents


def parse_text(file_path: str, start_id: int = 1) -> List[SourceDocument]:
    """Read plain text or Markdown files."""
    documents: List[SourceDocument] = []
    try:
        base_name = os.path.basename(file_path)
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read().strip()
        if content:
            documents.append(
                SourceDocument(
                    id=start_id,
                    title=f"{base_name}",
                    url_or_path=file_path,
                    snippet=content[:3000],
                    source_type="manual",
                    relevance_score=1.0,
                )
            )
    except Exception as exc:
        print(f"[OmniResearch Doc Warning] Error reading text file {file_path}: {exc}")
    return documents


def parse_document(file_path: str, start_id: int = 1) -> List[SourceDocument]:
    """Universal document dispatcher based on extension."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(file_path, start_id=start_id)
    elif ext in (".docx", ".doc"):
        return parse_docx(file_path, start_id=start_id)
    elif ext == ".csv":
        return parse_csv(file_path, start_id=start_id)
    elif ext in (".txt", ".md", ".json"):
        return parse_text(file_path, start_id=start_id)
    return []
