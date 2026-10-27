"""DOCX Word Exporter for OmniResearch Executive Dossiers."""

import os

from docx import Document
from docx.shared import Pt, RGBColor

from omniresearch.core.state import ResearchReport


def export_docx_file(report: ResearchReport, output_path: str) -> str:
    """Generate styled Microsoft Word (.docx) dossier."""
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc = Document()

    # Document Title
    title = doc.add_heading(level=0)
    run = title.add_run(f"Executive Dossier: {report.topic}")
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(24, 43, 73)

    # Subtitle metadata
    p_meta = doc.add_paragraph()
    p_meta.add_run(f"OmniResearch Intelligence Engine | {report.generated_at}\n").italic = True
    p_meta.add_run(f"Average Confidence Score: {report.average_confidence_score}%").bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 1. Executive Summary
    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(report.executive_summary)

    # 2. Key Findings
    if report.key_findings:
        doc.add_heading("2. Key Strategic Findings", level=1)
        for kf in report.key_findings:
            doc.add_paragraph(kf, style="List Bullet")

    # 3. Analytical Sections
    doc.add_heading("3. In-Depth Analysis", level=1)
    for idx, sec in enumerate(report.sections, start=1):
        doc.add_heading(f"3.{idx}. {sec.heading}", level=2)
        doc.add_paragraph(sec.content)

    # 4. Strategic Recommendations
    if report.strategic_recommendations:
        doc.add_heading("4. Strategic Recommendations", level=1)
        for idx, rec in enumerate(report.strategic_recommendations, start=1):
            doc.add_paragraph(f"{idx}. {rec}")

    # 5. Fact Check Matrix Table
    if report.fact_check_summary:
        doc.add_heading("5. Fact-Check & Verification Matrix", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Claim / Metric"
        hdr_cells[1].text = "Status"
        hdr_cells[2].text = "Confidence"
        hdr_cells[3].text = "Reference"

        for claim in report.fact_check_summary:
            row_cells = table.add_row().cells
            row_cells[0].text = claim.claim_text
            row_cells[1].text = claim.status.upper()
            row_cells[2].text = f"{claim.confidence_score}%"
            row_cells[3].text = f"Source {claim.source_id}" if claim.source_id else "Corroborated"

    # 6. References
    if report.sources:
        doc.add_heading("6. Primary Evidence & Cited Sources", level=1)
        for s in report.sources:
            p_src = doc.add_paragraph(style="List Bullet")
            p_src.add_run(f"[Source {s.id}] {s.title} ({s.source_type.upper()}): ").bold = True
            p_src.add_run(s.url_or_path).italic = True

    doc.save(output_path)
    return output_path
